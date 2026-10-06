"""Build the class and calibration tables that the web map reads.

Reads new_class.xlsx (23 field classes -> 17-class legend), color_code.txt
(legend colours), the TabPFN-3.5 calibration results and the raw prediction's
pixel counts from the repository, and writes:

    app/data/classes.json          for the backend
    app/data/calibration.json      for the backend
    public/data/tables.js          the same tables plus class pixel counts, for
                                   the page

Run it again whenever new_class.xlsx or color_code.txt changes:

    python scripts/build_tables.py
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

WEB_MAP_DIR = Path(__file__).resolve().parents[1]
REPO_DIR = WEB_MAP_DIR.parent
OUT_DIR = WEB_MAP_DIR / "public" / "data"
TABLES_DIR = WEB_MAP_DIR / "app" / "data"

RECLASS_XLSX = REPO_DIR / "new_class.xlsx"
COLOR_TABLE = REPO_DIR / "color_code.txt"
CALIBRATION_BINS = REPO_DIR / "outputs_tabpfn35_polygon" / "calibration_bins.csv"
CALIBRATION_SUMMARY = REPO_DIR / "outputs_tabpfn35_polygon" / "calibration_summary.json"
PIXEL_COUNTS = REPO_DIR / "predictions_tabpfn35_polygon" / "predicted_pixel_counts.csv"

TREE_CROP_NAME = "Tree Crop Plantation"
CALIBRATION_MODEL = "TabPFN-3.5"
# Every raw pixel of every held-out polygon, each polygon weighted equally:
# the level the README quotes for the confidence map.
CALIBRATION_LEVEL = "raw_pixel_polygon_weighted"

# Legend groups in display order, as in the README's area table.
GROUPS = {
    "Forest": ["Closed Forest", "Open Forest", "Mangrove"],
    "Savannah & shrubland": ["Closed Savannah", "Open Savannah", "Open Low Shrubland"],
    "Agriculture": ["Tree Crop Plantation", "Annual Herbaceous Cultivation", "Shrub Crop"],
    "Wetlands & water": ["Swamp", "Seasonally Flooded Woody", "Permanently Flooded Woody", "Water"],
    "Built-up & bare": ["Built up", "Mining", "Bare Rocks", "Beaches"],
}

# Readable labels for field classes whose names are abbreviations.
FIELD_LABELS = {
    "ANN": "Annual crops",
    "S_Flood": "Seasonally flooded woody",
    "P_Flood": "Permanently flooded woody",
    "Woodland Savanna": "Woodland savannah",
    "Open forest": "Open forest",
    "Open Low shrubland": "Open low shrubland",
    "Palm": "Oil palm",
}


def read_reclass_table() -> list[dict]:
    sheet = load_workbook(RECLASS_XLSX, read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    col = {name: header.index(name) for name in ("Old Class Code", "Old Class", "New Class", "New Class Code")}
    table = []
    for row in rows[1:]:
        if row[col["Old Class Code"]] is None:
            continue
        table.append({
            "old_code": int(row[col["Old Class Code"]]),
            "old_name": str(row[col["Old Class"]]).strip(),
            "new_name": str(row[col["New Class"]]).strip(),
            "new_code": int(row[col["New Class Code"]]),
        })
    return table


def read_colors() -> dict[int, dict]:
    colors = {}
    for line in COLOR_TABLE.read_text(encoding="utf-8").splitlines():
        parts = line.split()
        if len(parts) < 6:
            continue
        code, r, g, b = (int(v) for v in parts[:4])
        if code == 0:
            continue
        colors[code] = {"hex": f"#{r:02x}{g:02x}{b:02x}", "name": " ".join(parts[5:])}
    return colors


def build_classes() -> dict:
    table = read_reclass_table()
    colors = read_colors()

    legend: dict[int, dict] = {}
    for row in table:
        code = row["new_code"]
        if code not in colors:
            sys.exit(f"Legend class {code} ({row['new_name']}) has no colour in {COLOR_TABLE.name}.")
        if colors[code]["name"].lower() != row["new_name"].lower():
            print(f"Note: class {code} is '{row['new_name']}' in {RECLASS_XLSX.name} "
                  f"but '{colors[code]['name']}' in {COLOR_TABLE.name}; using the spreadsheet name.")
        entry = legend.setdefault(code, {"code": code, "name": row["new_name"], "color": colors[code]["hex"],
                                         "field_classes": []})
        entry["field_classes"].append(row["old_code"])

    group_of = {name: group for group, names in GROUPS.items() for name in names}
    order = {name: i for i, name in enumerate(n for names in GROUPS.values() for n in names)}
    for entry in legend.values():
        entry["group"] = group_of.get(entry["name"], "Other")
    ordered = sorted(legend.values(), key=lambda e: (order.get(e["name"], len(order)), e["code"]))

    tree = [e["code"] for e in ordered if e["name"] == TREE_CROP_NAME]
    if not tree:
        sys.exit(f"No legend class named '{TREE_CROP_NAME}' in {RECLASS_XLSX.name}.")

    field = sorted(table, key=lambda r: r["old_code"])
    return {
        "generated_from": [RECLASS_XLSX.name, COLOR_TABLE.name],
        "tree_crop_code": tree[0],
        "remap": {"from": [r["old_code"] for r in field], "to": [r["new_code"] for r in field]},
        "field_classes": [{"code": r["old_code"], "name": r["old_name"],
                           "label": FIELD_LABELS.get(r["old_name"], r["old_name"]),
                           "legend_code": r["new_code"]} for r in field],
        "legend": ordered,
        "groups": [g for g in GROUPS if any(e["group"] == g for e in ordered)]
                  + (["Other"] if any(e["group"] == "Other" for e in ordered) else []),
    }


def build_calibration() -> dict:
    def number(value: str):
        return float(value) if value not in ("", None) else None

    with CALIBRATION_BINS.open(newline="", encoding="utf-8") as f:
        rows = [r for r in csv.DictReader(f)
                if r["model"] == CALIBRATION_MODEL and r["level"] == CALIBRATION_LEVEL]
    summary_file = json.loads(CALIBRATION_SUMMARY.read_text(encoding="utf-8"))
    summary = summary_file["models"][CALIBRATION_MODEL][CALIBRATION_LEVEL]
    return {
        "model": CALIBRATION_MODEL,
        "level": CALIBRATION_LEVEL,
        "level_description": summary_file["levels"][CALIBRATION_LEVEL],
        "threshold": summary["threshold"],
        "accuracy_at_or_above": summary["accuracy_at_or_above"],
        "accuracy_below": summary["accuracy_below"],
        "share_below": summary["share_below"],
        "ece": summary["ece"],
        "bins": [{"low": number(r["bin_low"]), "high": number(r["bin_high"]),
                  "share": number(r["weight_share"]), "mean_confidence": number(r["mean_confidence"]),
                  "accuracy": number(r["accuracy"])} for r in rows],
    }


def build_pixel_counts(classes: dict) -> dict:
    """Pixel counts of the raw (unfiltered) 23-class prediction, summed into the legend."""
    to_legend = dict(zip(classes["remap"]["from"], classes["remap"]["to"]))
    counts: dict[int, int] = {}
    with PIXEL_COUNTS.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            code = int(row["code"])
            if code in to_legend:
                counts[to_legend[code]] = counts.get(to_legend[code], 0) + int(row["pixel_count"])
    return {
        "source": f"{PIXEL_COUNTS.parent.name}/{PIXEL_COUNTS.name}",
        "description": "Pixel counts of the raw 23-class prediction (no majority filter), summed into the 17-class "
                       "legend. Includes coastal sea pixels classified as water.",
        "total": sum(counts.values()),
        "by_class": {str(k): v for k, v in sorted(counts.items())},
    }


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TABLES_DIR.mkdir(parents=True, exist_ok=True)
    classes, calibration = build_classes(), build_calibration()
    for name, data in (("classes.json", classes), ("calibration.json", calibration)):
        (TABLES_DIR / name).write_text(json.dumps(data, indent=1), encoding="utf-8")
        print(f"Wrote {TABLES_DIR / name}")
    tables = {"classes": classes, "calibration": calibration, "pixel_counts": build_pixel_counts(classes)}
    (OUT_DIR / "tables.js").write_text(
        "// Generated by scripts/build_tables.py - do not edit by hand.\n"
        f"window.WM_TABLES = {json.dumps(tables, indent=1)};\n", encoding="utf-8")
    print(f"Wrote {OUT_DIR / 'tables.js'}")


if __name__ == "__main__":
    main()
